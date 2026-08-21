"""Turn whatever the taxpayer sent into a structure the completeness checker can read.

Deliberately separate from the checking. Extraction is format work — xlsx, csv, a PDF, a letter
in the body of an email — and it changes with the file. Checking is rule work and should not.
Keeping them apart is what lets the completeness checks stay deterministic no matter what
arrives, and lets a better extractor be dropped in without touching a single rule.

Column naming is the one place where being strict would produce nonsense. A taxpayer who writes
"Invoice No." has supplied the invoice number, and reporting it missing would be wrong. So
headers are normalised and run through a modest alias table. Anything the alias table does not
cover is reported as missing — and the document reviewer, which is allowed to exercise
judgement, can then observe that the column is present under a name we did not anticipate. The
deterministic layer stays strict; the judgement stays with the layer that is allowed to have it.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any

# canonical column name -> the spellings we accept for it
ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_date": ("date", "inv_date", "invoice_dt", "issue_date", "date_of_invoice",
                     "tax_invoice_date"),
    "invoice_number": ("invoice_no", "inv_no", "invoice_num", "invoice", "document_number",
                       "doc_no", "tax_invoice_number", "invoice_ref"),
    "customer_name": ("customer", "client_name", "client", "buyer_name", "buyer",
                      "customer_title"),
    "customer_vat_number": ("customer_vat", "customer_trn", "buyer_vat", "buyer_vat_number",
                            "customer_vat_no", "vat_number_of_customer"),
    "supplier_name": ("supplier", "vendor_name", "vendor", "seller_name", "seller"),
    "supplier_vat_number": ("supplier_vat", "vendor_vat", "seller_vat", "supplier_vat_no",
                            "supplier_trn"),
    "description": ("desc", "details", "narration", "particulars", "goods_description",
                    "line_description"),
    "taxable_amount": ("net_amount", "amount_excl_vat", "amount_before_vat", "taxable_value",
                       "net", "amount_excluding_vat", "base_amount"),
    # "vat" on its own is the amount far more often than the rate, so it belongs below
    "vat_rate": ("rate", "tax_rate", "vat_pct", "vat_percent", "vat_rate_pct"),
    "vat_amount": ("vat", "vat_value", "tax_amount", "vat_amt", "output_vat", "input_vat",
                   "vat_amount_sar"),
    "note_date": ("credit_note_date", "cn_date"),
    "note_number": ("credit_note_no", "cn_no", "note_no", "credit_note_number"),
    "note_type": ("type", "document_type", "doc_type"),
    "original_invoice_number": ("original_invoice", "related_invoice", "invoice_referenced",
                                "against_invoice"),
    "reason": ("reason_for_issue", "note_reason", "explanation"),
    "account_code": ("account", "gl_account", "account_no", "acc_code"),
    "account_name": ("account_description", "gl_name", "acc_name"),
    "debit": ("dr", "debit_amount"),
    "credit": ("cr", "credit_amount"),
    "posting_date": ("post_date", "gl_date", "transaction_date"),
    "reference": ("ref", "document_reference", "voucher"),
    "item": ("reconciling_item", "line", "particular"),
    "amount": ("value", "sar", "amount_sar"),
    "supporting_reference": ("support", "evidence", "reference_document"),
    "terminal_id": ("terminal", "pos_id", "device_id", "terminal_no"),
    "transactions": ("txn_count", "no_of_transactions", "count"),
    "gross_amount": ("gross", "total_amount", "amount_incl_vat"),
    "asset_code": ("asset_no", "asset_id", "fa_code"),
    "acquisition_date": ("purchase_date", "date_acquired", "addition_date"),
    "cost": ("acquisition_cost", "purchase_cost", "asset_cost"),
    "vat_claimed": ("input_vat_claimed", "vat_recovered"),
    "disposal_date": ("date_disposed", "disposal"),
    "date": ("day", "business_date", "trading_date"),
}

_ALIAS_LOOKUP: dict[str, str] = {}
for _canon, _spellings in ALIASES.items():
    _ALIAS_LOOKUP[_canon] = _canon
    for _s in _spellings:
        _ALIAS_LOOKUP.setdefault(_s, _canon)

TOTAL_ROW = re.compile(r"^\s*(grand\s+)?totals?\b", re.I)
NUMERIC = re.compile(r"^-?[\d,]+(\.\d+)?$")


def normalise(header: Any) -> str:
    """'Invoice No. ' -> 'invoice_no' -> canonical 'invoice_number'."""
    s = str(header or "").strip().lower()
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\s+", "_", s).strip("_")
    return _ALIAS_LOOKUP.get(s, s)


_UPPER = {"vat", "isic", "hs", "pos", "id", "no"}


def display(column: str) -> str:
    """'customer_vat_number' -> 'Customer VAT number'.

    Snake_case is the checker's vocabulary, not the taxpayer's. A letter that asks for
    'customer_vat_number' reads like a database error, so gap details are phrased with this.
    """
    words = str(column or "").split("_")
    if not words or not words[0]:
        return str(column or "")
    out = [words[0].upper() if words[0] in _UPPER else words[0].capitalize()]
    out += [w.upper() if w in _UPPER else w for w in words[1:]]
    return " ".join(out)


def _num(v: Any) -> float | None:
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v or "").strip().replace(",", "").replace("SAR", "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _as_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


# --------------------------------------------------------------------------- table assembly

def _from_grid(grid: list[list[Any]]) -> dict:
    """Find the header row, split off any stated totals, and normalise the columns."""
    header_idx = None
    for i, row in enumerate(grid[:20]):
        filled = [c for c in row if not _is_blank(c)]
        if len(filled) >= 2 and not all(_num(c) is not None for c in filled):
            header_idx = i
            break
    if header_idx is None:
        return {"columns": [], "rows": [], "stated_totals": {}, "raw_headers": []}

    raw = grid[header_idx]
    keep = [j for j, h in enumerate(raw) if not _is_blank(h)]
    raw_headers = [str(raw[j]).strip() for j in keep]
    columns = [normalise(raw[j]) for j in keep]

    rows: list[list[Any]] = []
    stated: dict[str, float] = {}
    for row in grid[header_idx + 1:]:
        cells = [row[j] if j < len(row) else None for j in keep]
        if all(_is_blank(c) for c in cells):
            continue
        label = next((str(c) for c in cells if isinstance(c, str) and c.strip()), "")
        if TOTAL_ROW.match(label):
            for col, cell in zip(columns, cells):
                n = _num(cell)
                if n is not None:
                    stated[col] = n
            continue
        rows.append(cells)
    return {"columns": columns, "rows": rows, "stated_totals": stated,
            "raw_headers": raw_headers}


def _period(columns: list[str], rows: list[list[Any]]) -> tuple[str | None, str | None]:
    """The span the document actually covers, from whichever column looks like a date."""
    idxs = [i for i, c in enumerate(columns) if c.endswith("date")]
    seen: list[date] = []
    for row in rows:
        for i in idxs:
            d = _as_date(row[i]) if i < len(row) else None
            if d:
                seen.append(d)
    if not seen:
        return None, None
    return min(seen).isoformat(), max(seen).isoformat()


# --------------------------------------------------------------------------- entry points

def extract(filename: str, data: bytes) -> dict:
    """Normalise one received file. Never raises — an unreadable file is a finding, not a crash."""
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    try:
        if ext in ("xlsx", "xlsm"):
            out = _extract_xlsx(data)
        elif ext in ("csv", "txt") and ext == "csv":
            out = _extract_csv(data)
        else:
            text = data.decode("utf-8", errors="replace")
            return {"format": ext or "unknown", "columns": [], "rows": [], "raw_headers": [],
                    "stated_totals": {}, "period_from": None, "period_to": None,
                    "text": text[:20000], "row_count": 0,
                    "note": "Not a tabular format — structured checks cannot run on it."}
    except Exception as exc:                       # noqa: BLE001 - report, do not crash
        return {"format": ext or "unknown", "columns": [], "rows": [], "raw_headers": [],
                "stated_totals": {}, "period_from": None, "period_to": None, "text": "",
                "row_count": 0, "note": f"Could not be read: {type(exc).__name__}."}

    pf, pt = _period(out["columns"], out["rows"])
    return {
        "format": ext, "columns": out["columns"], "raw_headers": out["raw_headers"],
        # keep rows JSON-safe: dates become strings
        "rows": [[c.isoformat() if isinstance(c, (date, datetime)) else c for c in r]
                 for r in out["rows"]],
        "stated_totals": out["stated_totals"], "period_from": pf, "period_to": pt,
        "text": "", "row_count": len(out["rows"]), "note": "",
    }


def _extract_xlsx(data: bytes) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return _from_grid(grid)


def _extract_csv(data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    grid = [row for row in csv.reader(io.StringIO(text))]
    return _from_grid(grid)


def from_structure(payload: dict) -> dict:
    """Accept an already-parsed table (demo seeding, or an upstream extraction service)."""
    columns = [normalise(c) for c in payload.get("columns", [])]
    rows = payload.get("rows", []) or []
    pf, pt = payload.get("period_from"), payload.get("period_to")
    if pf is None and pt is None:
        pf, pt = _period(columns, rows)
    return {
        "format": payload.get("format", "xlsx"), "columns": columns,
        "raw_headers": list(payload.get("columns", [])), "rows": rows,
        "stated_totals": payload.get("stated_totals", {}) or {},
        "period_from": pf, "period_to": pt, "text": payload.get("text", ""),
        "row_count": len(rows), "note": payload.get("note", ""),
    }
